#!/usr/bin/env python3
"""Build the Word project report from committed ER diagrams and screenshot evidence.
Requires python-docx. Run from any working directory.
"""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/delivery'
DIAGRAMS=ROOT/'docs/engineering/diagrams'
SCREENSHOTS=OUT/'screenshots'
REPORT=OUT/'邻里闲置_项目设计说明书.docx'

def east_font(style,name='Songti SC'):
    style.font.name='Times New Roman'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)

def paragraph(doc,content,style=None):
    p=doc.add_paragraph(content,style)
    p.paragraph_format.space_after=Pt(7)
    return p

def heading(doc,text,level=1):
    return doc.add_heading(text,level)

def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    widths=widths or ([4,13] if len(headers)==2 else [3,5,9])
    t.autofit=False
    for col,width in zip(t.columns,widths): col.width=Cm(width)
    for c,v in zip(t.rows[0].cells,headers):c.text=v
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,row):c.text=str(v)
    for i,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for j,cell in enumerate(row.cells):
            if widths: cell.width=Cm(widths[j])
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                for r in p.runs:
                    r.font.size=Pt(10);r.bold=(i==0)
                    if i==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF0EE');cell._tc.get_or_add_tcPr().append(shade)
    return t

def image(doc,path,caption,width=6):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(str(path),width=Inches(width))
    p=paragraph(doc,caption,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER

def new_page(doc,title):
    doc.add_page_break();heading(doc,title)

def main():
    OUT.mkdir(exist_ok=True)
    manifest_path=SCREENSHOTS/'manifest.json'
    evidence=json.loads(manifest_path.read_text()) if manifest_path.exists() else {'kind':'pending','pages':[]}
    d=Document();s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2);s.right_margin=Cm(2)
    s.header_distance=Cm(.8);s.footer_distance=Cm(.8)
    normal=d.styles['Normal'];east_font(normal);normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.15
    for name,size in [('Title',30),('Subtitle',15),('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
        st=d.styles[name];east_font(st,'PingFang SC');st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('18372E')
    east_font(d.styles['Caption']);d.styles['Caption'].font.size=Pt(10)
    h=s.header.paragraphs[0];h.text='邻里闲置  |  项目设计说明书';h.style='Caption'
    f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f.add_run('第 ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    f.add_run(' 页')
    d.core_properties.title='邻里闲置 · 项目设计说明书'
    d.core_properties.subject='产品说明、Web 页面截图、Chen E-R 图与系统设计'
    d.core_properties.author='邻里闲置项目组'
    for _ in range(4):paragraph(d,'')
    p=d.add_paragraph('邻里闲置','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph('项目设计说明书','Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=paragraph(d,'社区闲置物品流转 Web 应用');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    paragraph(d,'')
    table(d,['项目','说明'],[
        ('应用形态','移动端优先、兼容桌面浏览器的 Web 应用'),
        ('文档内容','产品说明、Web 页面截图、技术方案、10 张独立 Chen E-R 图'),
        ('文档阶段','设计阶段；应用业务与数据库尚未实现'),
        ('截图类型','本地 Web 页面原型截图，使用虚构数据' if evidence['kind']=='prototype' else '实际运行截图待 Web 功能完成后补入'),
        ('版本 / 日期','1.0 / 2026-09-27')])
    paragraph(d,'本报告用于说明产品与数据库设计。实现、真实 LLM 调用、运行测试和部署结果需在开发完成后补充核验；设计图或页面原型不作为全栈验收证据。')
    new_page(d,'1  项目概述与需求')
    heading(d,'1.1 场景与目标',2)
    paragraph(d,'小区群里的闲置物品容易被新消息淹没；领取者难以判断物品是否还在，双方又需要反复私聊约时间和地点。邻里闲置面向同一个小区、楼栋或办公室，以物品卡片、明确状态和预约交接减少这些沟通成本。')
    paragraph(d,'本次以单一预置社区为范围，采用虚构演示身份和预置图片，优先完成一条可验证的线下交易闭环。')
    heading(d,'1.2 主要用户与流程',2)
    paragraph(d,'发布者与领取者都是社区用户，同一用户可在不同物品下承担不同角色。访客可以浏览；写操作由服务端校验身份和权限。')
    paragraph(d,'发布物品 → 浏览与搜索 → 表达想要 → 发布者选人并预约 → 领取者确认 → 线下交接 → 发布者确认送出 → 归档与统计更新。')
    table(d,['功能','设计说明'],[
        ('物品发布','名称、描述、交易方式、预置图片与自提楼栋；支持免费送、随便给、标价。'),
        ('发现与检索','卡片按发布时间倒序；支持关键词搜索和交易方式筛选；显示新鲜度。'),
        ('意向与交接','想要、撤回、查看意向人；约定未来 7 天内的时间和公共地点；确认、取消与完成。'),
        ('留言与归档','公开问答减少重复私聊；已送出保留记录并禁止写操作。'),
        ('社区看板','本月发布、本月成交、当前在售、最快领走和最想要物品。'),
        ('LLM 辅助','根据已知事实整理文案和给出参考价；用户选择采用，失败时仍能手动发布。')])
    heading(d,'1.3 取舍与交付',2)
    paragraph(d,'不纳入本次基线：支付、快递、即时私聊、信用体系、多社区、真实注册和图片上传。考试交付包括本地可运行全栈应用、5 分钟视频、数据库 E-R 图及技术栈简介；远端部署为可选项。')
    new_page(d,'2  技术架构与系统规则')
    table(d,['层次','选型','职责'],[
        ('前端','React + TypeScript + Vite','页面、表单、状态反馈与 API 调用'),
        ('后端','Node.js + Fastify','身份、权限、参数校验、业务事务与 LLM 接入'),
        ('数据库','SQLite + Drizzle','物品、意向、留言、交易和会话的持久化'),
        ('接口契约','OpenAPI 3.1 + 共享 Schema','固定输入、输出与错误结构'),
        ('验证','Vitest + API 测试 + Playwright','业务规则、权限、并发与双会话端到端验收')])
    paragraph(d,'浏览器通过同源 /api/v1 访问后端；后端访问 SQLite 并调用模型服务。密钥只放在服务端环境变量中。前端、API 与数据库保持明确边界。上述选型为计划，依赖版本将在工程初始化时锁定。')
    heading(d,'2.1 状态与一致性',2)
    paragraph(d,'物品：可领取 → 已预约 → 已送出；取消预约后恢复可领取。交易：待确认 → 已确认 → 已完成；待确认或已确认的交易允许双方取消。物主不能领取自己的物品，只有指定领取者能确认预约，只有物主能确认送出。')
    paragraph(d,'预约创建、取消和完成在事务内校验状态并更新数据。部分唯一索引保证同一物品最多一笔未取消交易；重复确认或完成不会重复成交。交易细节只对双方开放。')
    heading(d,'2.2 时间、价格与统计',2)
    paragraph(d,'不足 24 小时为“刚上架”；24 至不足 72 小时为“新上架”；之后按经过整天数展示。金额以整数分保存。月统计按北京时间计算，成交仅统计已完成交易，当前在售含可领取和已预约。')
    heading(d,'2.3 模型能力的边界',2)
    paragraph(d,'AI 只根据输入事实优化描述，不擅自补造品牌、成色或故障状况。参考价注明不是行情估价；响应需结构校验，超时或无配置时保留用户输入。最终验收须单独记录一次真实模型调用。')
    if evidence['pages']:
        for i,page in enumerate(evidence['pages'],1):
            new_page(d,f'3.{i}  Web 页面：{page["title"]}')
            paragraph(d,'截图来源：本地浏览器打开项目 Web 原型后实拍。页面中的数据与 AI 建议为示例；此原型用于说明交互设计，未连接后端、数据库或真实模型。')
            image(d,SCREENSHOTS/page['file'],f'图 3-{i}  {page["title"]}（Web 原型截图）',6.3)
            heading(d,'页面说明',2);paragraph(d,page['description'])
            paragraph(d,f'采集条件：{page["viewport"]} 浏览器视口；对应原型页面：{page["route"]}。')
    else:
        new_page(d,'3  Web 页面截图')
        paragraph(d,'当前尚无可运行 Web 应用，本章等待开发完成后采集实际运行截图。下表明确每张图需要展示的行为；不使用设计图代替运行证据。')
        table(d,['页面','截图说明'],[('社区首页','物品卡片、新鲜度、搜索筛选与统计'),('物品详情','当前状态、物品信息、想要与留言'),('发布页面','发布表单与 AI 建议的采用方式'),('我的交接','意向用户、时间地点、确认或取消与完成')])
    new_page(d,'4  数据库概念结构与全局 E-R 图')
    paragraph(d,'采用 Chen 表示法：矩形表示实体，椭圆表示属性，菱形表示联系，主键属性加下划线，直线无箭头，1 / n 表示联系的最大基数。共包含 6 张实体属性图、3 张局部联系图和 1 张全局图。')
    image(d,DIAGRAMS/'er-global.png','图 4-1  全局实体联系图',6.0)
    paragraph(d,'六个实体为用户、登录会话、物品、意向记录、留言和交易记录。取消联系可选，交易有零个或一个取消人；物品与交易的一对多包含取消历史，不表示允许同时成交给多人。')
    attrs=[('user','用户'),('session','登录会话'),('item','物品'),('interest','意向记录'),('comment','留言'),('trade','交易记录')]
    for group in range(3):
        new_page(d,f'5.{group+1}  实体属性图')
        for j in range(2):
            i=group*2+j;slug,name=attrs[i]
            image(d,DIAGRAMS/f'attribute-{slug}.png',f'图 5-{i+1}  {name}实体属性图',5.65)
        if group==0:paragraph(d,'标识属性以椭圆内文字下划线表示。外键通过联系图表达；字段英文名、数据类型和约束见 SRS。')
    locals=[('publishing','用户发布与登录','一个用户可拥有多个会话、发布多件物品。会话与物品各属于一个用户。'),('interest-comments','物品意向与留言','每条意向或留言关联一个用户和一件物品。同一用户对同一物品只有一条意向记录，撤回后可再次激活。'),('trading','交易交接','每笔交易关联一件物品和一个指定领取者；发布者由物品确定。取消人必须是双方之一；同一物品最多一笔未取消交易。')]
    for i,(slug,name,note) in enumerate(locals,1):
        new_page(d,f'6.{i}  局部实体联系图：{name}')
        image(d,DIAGRAMS/f'local-{slug}.png',f'图 6-{i}  {name}联系图',6.3)
        paragraph(d,note)
    new_page(d,'7  关系模式与完整性约束')
    table(d,['关系表','主键与关键引用','主要约束'],[
        ('users','id','昵称与楼栋；用户可同时承担发布和领取角色'),
        ('sessions','id；user_id → users','token_hash 唯一，仅保存令牌哈希，支持过期'),
        ('items','id；owner_id → users','交易方式与金额匹配；GIVEN 必须有送出时间'),
        ('interests','id；item_id、user_id','(item_id,user_id) 联合唯一；active 标记有效状态'),
        ('comments','id；item_id、author_id','归档物品不可新增留言；正文为纯文本'),
        ('trades','id；item_id、recipient_id、cancelled_by','状态时间字段一致；取消人可空；未取消交易部分唯一')])
    paragraph(d,'ID 使用 UUID；数据库时间采用 UTC 毫秒，API 使用带时区的 ISO 8601。外键、CHECK 与唯一索引由迁移建立；跨表身份和状态条件在同一事务中校验。图纸在开发完成后还须对照实际迁移复核。')
    heading(d,'7.1 关键约束示例',2)
    paragraph(d,'免费送价格为 0；随便给无固定价格；标价为正整数分。一个物品可以有多次取消的历史交易，但 PENDING、CONFIRMED、COMPLETED 范围内最多一条交易。')
    paragraph(d,'交易完成时，交易状态与物品状态同时更新，完成时间保持一致。重复完成返回原结果，不新增成交记录；旧取消请求不能影响后续新预约。')
    new_page(d,'8  开发计划、验证与交付状态')
    paragraph(d,'GitHub Epic #1 下设 11 个必需开发任务，另有 1 个可选部署任务。先建立工程与契约，再实施数据库和模块，最后联调、验证并制作交付材料。页面原型与本报告不代表这些开发任务已完成。')
    table(d,['验证项','验收要求','当前状态'],[
        ('双角色交易','发布→想要→预约→确认→完成→归档','待实现与验证'),
        ('权限与并发','越权拒绝、同一物品唯一有效预约、重试幂等','待实现与验证'),
        ('数据持久化','迁移、种子及服务重启后数据保存','待实现与验证'),
        ('真实 LLM','成功调用证据与超时/缺配置降级','待配置与验证'),
        ('Web 截图','截图来源、页面用途和演示数据标注','原型截图已纳入' if evidence['pages'] else '等待实际运行截图'),
        ('E-R 图','10 张独立 Chen 图已绘制，后续核对实际迁移','设计图已完成'),
        ('5 分钟视频','可运行 Demo 和实现讲解','待开发完成后录制'),
        ('CI / 部署','配置后跟踪终态，部署后核验线上','当前未配置')])
    heading(d,'8.1 配套资料',2)
    paragraph(d,'产品范围：docs/product/PRD.md；系统规格：docs/engineering/SRS.md；图纸与映射：docs/engineering/ERD.md；依赖计划：docs/planning/DAG.md。')
    paragraph(d,'仓库：https://github.com/luochen211/neighborhood-exchange\n交付 Epic：https://github.com/luochen211/neighborhood-exchange/issues/1')
    paragraph(d,'后续报告应逐项记录通过、实际失败、跳过与外部阻塞；不得把原型、模拟模型响应或计划中的行为写成已完成的全栈测试结果。')
    d.save(REPORT)
    print(REPORT)

if __name__=='__main__':main()
